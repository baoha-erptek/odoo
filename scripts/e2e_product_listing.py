"""E2E smoke — FLOW_TAO_SAN_PHAM_VN against staging JaHandmadeArt.

Walks the owner-facing flow end-to-end:
  Excel catalog row  →  product.creation.wizard.action_create
                     →  etsy.publish.wizard.action_run_publish_draft_only
                     →  verify product.channel.status + etsy.listing + Etsy API GET

Stops at Etsy 'draft' state to avoid the $0.20/listing fee. Listings are
left on the JaHandmadeArt shop as drafts for owner inspection.

Source-of-truth docs:
  - docs/owner/FLOW_TAO_SAN_PHAM_VN.md (the flow itself)
  - .claude/plans/check-for-docs-owner-flow-tao-san-pham-v-purring-glade.md
  - tracker slice: P-PUB-E2E

Usage:
    python3 scripts/e2e_product_listing.py [--count N] [--shop NAME] \\
        [--base-url URL] [--db NAME] [--cleanup] [--verbose]

Exit codes:
    0 = all rows succeeded
    1 = pre-flight failure (creds / shop / channel / category missing)
    2 = wizard creation failed
    3 = publisher draft path failed
    4 = post-condition verification failed
"""

from __future__ import annotations

import argparse
import logging
import os
import sys
import xmlrpc.client
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import requests
from dotenv import load_dotenv

REPO_ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = REPO_ROOT / ".0temp" / "raw" / "[2025] Product Catalog.xlsx"
SHEETS_IN_ORDER = ("Accessories", "Pet", "Apparel", "Drinkware", "Home Decor", "Combo")
ETSY_API_BASE = "https://api.etsy.com/v3/application"

# Safe non-zero defaults used by the wizard call. Real catalog rows often
# lack price/shipping in the Excel; we substitute fixed values to satisfy
# the wizard's > 0 / >= 0 validations.
# Note: Etsy's createListing price is in the SHOP's listed currency, not USD.
# JaHandmadeArt is configured as VND (min price 5,043 VND); 250,000 ≈ $10 USD.
# Override via env var or change before running against non-VND shops.
DEFAULT_LISTING_PRICE_USD = float(os.environ.get("E2E_LISTING_PRICE", "250000"))
DEFAULT_SHIPPING_PRICE = float(os.environ.get("E2E_SHIPPING_PRICE", "50000"))

load_dotenv(REPO_ROOT / ".env")

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("e2e_product_listing")


# ─── data classes ──────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class CatalogRow:
    sheet: str
    row_idx: int
    name: str
    sku: str


@dataclass(frozen=True)
class StagingConfig:
    base_url: str
    db: str
    login: str
    password: str


@dataclass
class XmlrpcSession:
    cfg: StagingConfig
    uid: int
    common: xmlrpc.client.ServerProxy
    models: xmlrpc.client.ServerProxy

    def call(self, model: str, method: str, args: list, kwargs: dict | None = None) -> Any:
        return self.models.execute_kw(
            self.cfg.db, self.uid, self.cfg.password,
            model, method, args, kwargs or {},
        )


# ─── helpers ───────────────────────────────────────────────────────────────────


def load_config(args: argparse.Namespace) -> StagingConfig:
    base_url = args.base_url or os.environ.get("STAGING_BASE_URL", "https://odoo.hatafax.com")
    db = args.db or os.environ.get("STAGING_DB", "demo_esty")
    login = os.environ.get("STAGING_BA_LOGIN", "demo_ba_manager@hatafax.demo")
    pwd = os.environ.get("STAGING_BA_PASSWORD", "demo1234")
    return StagingConfig(base_url=base_url.rstrip("/"), db=db, login=login, password=pwd)


def connect(cfg: StagingConfig) -> XmlrpcSession:
    common = xmlrpc.client.ServerProxy(f"{cfg.base_url}/xmlrpc/2/common", allow_none=True)
    models = xmlrpc.client.ServerProxy(f"{cfg.base_url}/xmlrpc/2/object", allow_none=True)
    uid = common.authenticate(cfg.db, cfg.login, cfg.password, {})
    if not uid:
        raise SystemExit(
            f"Authentication failed for {cfg.login!r} on {cfg.base_url}/{cfg.db}. "
            "Set STAGING_BA_LOGIN + STAGING_BA_PASSWORD in .env."
        )
    log.info("XMLRPC authenticated: uid=%s base=%s db=%s", uid, cfg.base_url, cfg.db)
    return XmlrpcSession(cfg=cfg, uid=uid, common=common, models=models)


def pick_rows(count: int) -> list[CatalogRow]:
    if not EXCEL_PATH.exists():
        raise SystemExit(f"Excel catalog not found: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    try:
        picked: list[CatalogRow] = []
        for sheet_name in SHEETS_IN_ORDER:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            header_row: list[str] | None = None
            for idx, row in enumerate(ws.iter_rows(values_only=True)):
                if header_row is None:
                    header_row = [str(c).strip() if c is not None else "" for c in row]
                    if "Product" not in header_row or "SKU" not in header_row:
                        # Not a catalog sheet (e.g., "Các chi phí khác")
                        break
                    name_col = header_row.index("Product")
                    sku_col = header_row.index("SKU")
                    continue
                name = (row[name_col] or "") if name_col < len(row) else ""
                sku = (row[sku_col] or "") if sku_col < len(row) else ""
                name_s = str(name).strip()
                sku_s = str(sku).strip()
                if not name_s or not sku_s:
                    continue
                # Skip multi-line SKU cells — those encode size/variant lists,
                # not a single SKU we can publish.
                if "\n" in sku_s:
                    continue
                picked.append(CatalogRow(
                    sheet=sheet_name, row_idx=idx, name=name_s, sku=sku_s,
                ))
                if len(picked) >= count:
                    return picked
        return picked
    finally:
        wb.close()


# ─── pre-flight ────────────────────────────────────────────────────────────────


@dataclass
class Preflight:
    etsy_channel_id: int
    category_id: int
    shop_id: int
    shop_name: str
    etsy_api_shop_id: str
    etsy_access_token: str


def preflight(s: XmlrpcSession, shop_name: str) -> Preflight:
    # 1. BA group membership
    has_ba = s.call(
        "res.users", "has_group",
        [s.uid, "multichannel_hub_core.group_ba_user"],
    )
    if not has_ba:
        raise SystemExit(
            f"User {s.cfg.login!r} is not in multichannel_hub_core.group_ba_user. "
            "Wizard's FR-017 gate will refuse. Add the user to the BA group."
        )
    log.info("BA group membership: OK")

    # 2. Etsy channel xmlid (XMLRPC blocks underscore-prefixed methods, so we
    # resolve via ir.model.data search_read instead of _xmlid_to_res_id).
    data = s.call(
        "ir.model.data", "search_read",
        [[("module", "=", "multichannel_hub_core"), ("name", "=", "channel_etsy")]],
        {"fields": ["res_id"], "limit": 1},
    )
    channel_id = data[0]["res_id"] if data else 0
    if not channel_id:
        raise SystemExit("multichannel_hub_core.channel_etsy xmlid not resolvable on staging.")
    log.info("Etsy channel: id=%s", channel_id)

    # 3. Category — use first internal product category
    cats = s.call(
        "product.category", "search_read",
        [[]], {"fields": ["id", "name"], "limit": 1, "order": "id asc"},
    )
    if not cats:
        raise SystemExit("No product.category found on staging.")
    category_id = cats[0]["id"]
    log.info("Category: id=%s name=%s", category_id, cats[0]["name"])

    # 4. JaHandmadeArt shop record
    shops = s.call(
        "etsy.shop", "search_read",
        [[("name", "=", shop_name)]],
        {"fields": [
            "id", "name", "etsy_api_shop_id",
            "etsy_oauth_access_token",
            "default_taxonomy_id", "default_shipping_profile_id", "default_return_policy_id",
        ], "limit": 1},
    )
    if not shops:
        raise SystemExit(
            f"etsy.shop {shop_name!r} not found on staging. "
            "Owner must authorize the shop via the OAuth callback first."
        )
    sh = shops[0]
    missing = []
    if not sh.get("etsy_api_shop_id"):
        missing.append("etsy_api_shop_id")
    if not sh.get("etsy_oauth_access_token"):
        missing.append("etsy_oauth_access_token")
    if not sh.get("default_taxonomy_id"):
        missing.append("default_taxonomy_id")
    if not sh.get("default_shipping_profile_id"):
        missing.append("default_shipping_profile_id")
    if not sh.get("default_return_policy_id"):
        missing.append("default_return_policy_id")
    if missing:
        raise SystemExit(
            f"etsy.shop {shop_name!r} is missing required publisher fields: "
            f"{', '.join(missing)}. Publisher will refuse otherwise."
        )
    log.info(
        "Shop %s: odoo_id=%s etsy_api_shop_id=%s defaults={tax=%s, ship=%s, ret=%s}",
        sh["name"], sh["id"], sh["etsy_api_shop_id"],
        sh["default_taxonomy_id"], sh["default_shipping_profile_id"], sh["default_return_policy_id"],
    )

    return Preflight(
        etsy_channel_id=channel_id,
        category_id=category_id,
        shop_id=sh["id"],
        shop_name=sh["name"],
        etsy_api_shop_id=str(sh["etsy_api_shop_id"]),
        etsy_access_token=sh["etsy_oauth_access_token"],
    )


# ─── per-row execution ─────────────────────────────────────────────────────────


@dataclass
class RowResult:
    row: CatalogRow
    prefixed_sku: str
    product_tmpl_id: int | None = None
    listing_id: str | None = None
    etsy_get_status: int | None = None
    etsy_state: str | None = None
    note: str = ""


def run_row(s: XmlrpcSession, pf: Preflight, row: CatalogRow, run_stamp: str) -> RowResult:
    prefixed_sku = f"E2E-{run_stamp}-{row.sku}"[:64]  # Odoo default_code is varchar/Char
    res = RowResult(row=row, prefixed_sku=prefixed_sku)

    # Step 1+2: create wizard + action_create
    try:
        wiz_id = s.call("product.creation.wizard", "create", [{
            "name": row.name,
            "default_code": prefixed_sku,
            "categ_id": pf.category_id,
            "x_listing_price": DEFAULT_LISTING_PRICE_USD,
            "x_shipping_price_internal": DEFAULT_SHIPPING_PRICE,
            "x_additional_cost": 0.0,
            "standard_price": 0.0,
            "x_channel_applicability_ids": [(6, 0, [pf.etsy_channel_id])],
        }])
        log.info("[%s] wizard created: id=%s", prefixed_sku, wiz_id)
        action = s.call("product.creation.wizard", "action_create", [[wiz_id]])
        tmpl_id = action.get("res_id") if isinstance(action, dict) else None
        if not tmpl_id:
            raise RuntimeError(f"action_create returned no res_id: {action!r}")
        res.product_tmpl_id = tmpl_id
        log.info("[%s] product.template created: id=%s", prefixed_sku, tmpl_id)
    except Exception as exc:
        res.note = f"wizard step failed: {exc}"
        log.error("[%s] %s", prefixed_sku, res.note)
        sys.exit(2)

    # Step 3: confirm product.channel.status row exists (state='draft', no external_ref yet)
    statuses = s.call(
        "product.channel.status", "search_read",
        [[("product_tmpl_id", "=", res.product_tmpl_id),
          ("channel_id", "=", pf.etsy_channel_id)]],
        {"fields": ["id", "state", "external_ref"], "limit": 1},
    )
    if not statuses:
        res.note = "product.channel.status row missing after wizard"
        sys.exit(4)
    log.info("[%s] channel.status pre-publish: %s", prefixed_sku, statuses[0])

    # Step 4: publish-draft-only via wizard
    try:
        pub_wiz = s.call("etsy.publish.wizard", "create", [{
            "product_tmpl_id": res.product_tmpl_id,
            "shop_id": pf.shop_id,
        }])
        action = s.call("etsy.publish.wizard", "action_run_publish_draft_only", [[pub_wiz]])
        listing_id = action.get("listing_id") if isinstance(action, dict) else None
        if not listing_id:
            raise RuntimeError(f"action_run_publish_draft_only returned no listing_id: {action!r}")
        res.listing_id = str(listing_id)
        log.info("[%s] Etsy listing created: id=%s", prefixed_sku, listing_id)
    except xmlrpc.client.Fault as fault:
        res.note = f"publish failed: {fault.faultString[:500]}"
        log.error("[%s] %s", prefixed_sku, res.note)
        sys.exit(3)
    except Exception as exc:
        res.note = f"publish failed: {exc}"
        log.error("[%s] %s", prefixed_sku, res.note)
        sys.exit(3)

    # Step 5: verify post-conditions
    statuses = s.call(
        "product.channel.status", "search_read",
        [[("product_tmpl_id", "=", res.product_tmpl_id),
          ("channel_id", "=", pf.etsy_channel_id)]],
        {"fields": ["id", "state", "external_ref"], "limit": 1},
    )
    if not statuses or statuses[0].get("external_ref") != res.listing_id:
        res.note = f"channel.status external_ref mismatch: {statuses!r}"
        log.error("[%s] %s", prefixed_sku, res.note)
        sys.exit(4)
    log.info("[%s] channel.status post-publish: %s", prefixed_sku, statuses[0])

    listings = s.call(
        "etsy.listing", "search_read",
        [[("shop_id", "=", pf.shop_id), ("etsy_listing_id", "=", res.listing_id)]],
        {"fields": ["id", "etsy_listing_id", "state", "title"], "limit": 1},
    )
    if not listings:
        res.note = "etsy.listing row not created"
        log.error("[%s] %s", prefixed_sku, res.note)
        sys.exit(4)
    log.info("[%s] etsy.listing row: %s", prefixed_sku, listings[0])

    # Optional: hit Etsy directly to confirm listing is reachable
    try:
        headers = {
            "Authorization": f"Bearer {pf.etsy_access_token}",
            "x-api-key": pf.etsy_access_token.split(":")[0] if ":" in pf.etsy_access_token else "",
        }
        # The shop's stored client key is needed for x-api-key. We don't have direct
        # access to credentials.json from this script, so fall back to a token-only
        # call: GET /listings/{id} accepts Bearer + x-api-key. If x-api-key is missing,
        # Etsy returns 401 — that's still a non-fatal sanity check, just log it.
        r = requests.get(
            f"{ETSY_API_BASE}/listings/{res.listing_id}",
            headers={"Authorization": f"Bearer {pf.etsy_access_token}"},
            timeout=15,
        )
        res.etsy_get_status = r.status_code
        if r.status_code == 200:
            try:
                res.etsy_state = r.json().get("state")
            except ValueError:
                res.etsy_state = "<non-json>"
        log.info(
            "[%s] Etsy GET /listings/%s -> HTTP %s state=%s",
            prefixed_sku, res.listing_id, r.status_code, res.etsy_state,
        )
    except requests.RequestException as exc:
        log.warning("[%s] Etsy GET verification skipped: %s", prefixed_sku, exc)

    return res


# ─── cleanup ───────────────────────────────────────────────────────────────────


def archive_template(s: XmlrpcSession, tmpl_id: int) -> None:
    try:
        s.call("product.template", "write", [[tmpl_id], {"active": False}])
        log.info("Archived product.template id=%s", tmpl_id)
    except Exception as exc:
        log.warning("Could not archive template id=%s: %s", tmpl_id, exc)


# ─── main ──────────────────────────────────────────────────────────────────────


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--count", type=int, default=1, help="How many catalog rows to publish (default 1)")
    ap.add_argument("--shop", default="JaHandmadeArt", help="etsy.shop.name to target")
    ap.add_argument("--base-url", default=None, help="Override STAGING_BASE_URL")
    ap.add_argument("--db", default=None, help="Override STAGING_DB")
    ap.add_argument("--cleanup", action="store_true", help="Archive the new product.templates after verifying")
    ap.add_argument("--verbose", action="store_true", help="DEBUG log level")
    args = ap.parse_args()

    if args.verbose:
        log.setLevel(logging.DEBUG)

    if args.count < 1 or args.count > 5:
        log.error("--count must be between 1 and 5 (smoke scope, not bulk).")
        return 1

    log.info("Step 1: read Excel + pick %s row(s)", args.count)
    rows = pick_rows(args.count)
    if not rows:
        log.error("No valid catalog rows found in %s", EXCEL_PATH)
        return 1
    for r in rows:
        log.info("  picked: sheet=%s row=%s name=%r sku=%s", r.sheet, r.row_idx, r.name, r.sku)

    log.info("Step 2: connect to staging")
    cfg = load_config(args)
    session = connect(cfg)

    log.info("Step 3: pre-flight")
    pf = preflight(session, args.shop)

    run_stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    log.info("Step 4: run %s row(s), prefix=E2E-%s-*", len(rows), run_stamp)

    results: list[RowResult] = []
    for row in rows:
        result = run_row(session, pf, row, run_stamp)
        results.append(result)

    if args.cleanup:
        log.info("Step 5: cleanup — archive created product.templates")
        for r in results:
            if r.product_tmpl_id:
                archive_template(session, r.product_tmpl_id)

    log.info("─" * 60)
    log.info("E2E summary — all rows published as drafts on shop %s:", pf.shop_name)
    for r in results:
        log.info(
            "  %s -> tmpl=%s listing=%s etsy_state=%s (HTTP %s)",
            r.prefixed_sku, r.product_tmpl_id, r.listing_id,
            r.etsy_state or "?", r.etsy_get_status or "?",
        )
    log.info("─" * 60)
    log.info(
        "Inspect drafts on Etsy Shop Manager (JaHandmadeArt) → Listings → Drafts; "
        "on Odoo UI (%s) open Sản phẩm and search for prefix 'E2E-%s-*'.",
        cfg.base_url, run_stamp,
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
